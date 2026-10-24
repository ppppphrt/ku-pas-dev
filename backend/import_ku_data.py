"""Read the supplied KU workbooks without loading student enrollment records."""

import argparse
import gzip
import hashlib
import json
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

from course_data import normalize
from schemas import Course, Section, TimeSlot

SOURCE_URL = "https://kudrive.ku.ac.th/index.php/s/s8ccWBcJLGAGDaT"
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def read_rows(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        fields = next(rows)
        for row in rows:
            yield dict(zip(fields, row, strict=True))
    finally:
        workbook.close()


def clock_time(value):
    # The export uses minutes since midnight, not HHMM.
    if not isinstance(value, (int, float)) or int(value) != value or not 0 <= value < 1440:
        raise ValueError("Invalid minute of day")
    return f"{int(value) // 60:02d}:{int(value) % 60:02d}"


def import_data(source: Path, destination: Path):
    destination.mkdir(parents=True, exist_ok=True)
    hashes = {
        name: hashlib.sha256((source / f"{name}.xlsx").read_bytes()).hexdigest()
        for name in ["subjects", "sections", "schedules"]
    }
    slots = defaultdict(list)
    invalid_sections = set()
    errors = []
    counts = Counter()
    for number, row in enumerate(read_rows(source / "schedules.xlsx"), 2):
        counts["schedule_rows"] += 1
        key = str(row["section_id"])
        try:
            start, end = clock_time(row["time_from"]), clock_time(row["time_to"])
            days = [day for day in DAYS if str(row[day.lower()]) == "1"]
            if not days or start >= end:
                raise ValueError("Missing day or nonpositive interval")
            for day in days:
                slot = TimeSlot(day=day, start=start, end=end).model_dump()
                if slot not in slots[key]:
                    slots[key].append(slot)
        except ValueError as exc:
            invalid_sections.add(key)
            errors.append({"record": number, "section_id": key, "reason": str(exc)})
    sections = defaultdict(list)
    section_ids = set()
    for row in read_rows(source / "sections.xlsx"):
        counts["section_rows"] += 1
        key = str(row["section_id"])
        if key in section_ids:
            raise ValueError(f"Duplicate section identity {key}")
        section_ids.add(key)
        times = sorted(slots.get(key, []), key=lambda x: (DAYS.index(x["day"]), x["start"]))
        complete = bool(times) and key not in invalid_sections
        note = (
            None if complete else "Schedule is missing or invalid in the source; cannot add safely."
        )
        section = Section(
            section_id=key,
            section_code=str(row["section_code"]),
            faculty=normalize(row["faculty_name_th"]),
            campus=normalize(row["campus_name_th"]),
            time_slots=times,
            schedule_complete=complete,
            schedule_note=note,
            total_seat=int(row["total_seat"]),
            total_registered=int(row["total_registered"]),
        )
        sections[normalize(row["subject_code"])].append(section.model_dump())
    courses = []
    codes = set()
    for row in read_rows(source / "subjects.xlsx"):
        counts["subject_rows"] += 1
        code = normalize(row["subject_code"])
        if code in codes:
            raise ValueError(f"Duplicate subject identity {code}; resolve before joining")
        codes.add(code)
        course = Course(
            subject_code=code,
            **{
                f: normalize(row[f])
                for f in [
                    "subject_name_th",
                    "subject_name_en",
                    "subject_desc_th",
                    "subject_desc_en",
                ]
            },
            credits=row["credit"],
            faculty=normalize(row["faculty_name_th"]),
            campus=normalize(row["campus_name_th"]),
            sections=sections.get(code, []),
        )
        courses.append(course.model_dump())
    orphan_codes = sorted(set(sections) - codes)
    orphan_schedule_ids = sorted(set(slots) - section_ids)
    counts["courses_with_sections"] = sum(bool(c["sections"]) for c in courses)
    counts["joinable_sections"] = sum(len(c["sections"]) for c in courses)
    counts["schedulable_sections"] = sum(
        s["schedule_complete"] for c in courses for s in c["sections"]
    )
    counts["invalid_schedule_rows"] = len(errors)
    counts["sections_without_schedule"] = len(section_ids - set(slots))
    counts["missing_credits"] = sum(c["credits"] is None for c in courses)
    content = json.dumps(courses, ensure_ascii=False, separators=(",", ":"))
    version = "ku-workbooks-v1:" + hashlib.sha256(content.encode()).hexdigest()
    metadata = {
        "dataset_version": version,
        "updated_at": datetime.now(UTC).isoformat(),
        "source_url": SOURCE_URL,
        "source_sha256": hashes,
        "academic_year": None,
        "semester": None,
        "period_note": "Academic year and semester are not present in the supplied export.",
        "data_source": "KU Drive workbook snapshot",
        "counts": dict(counts),
    }
    (destination / "catalog.json").write_text(
        json.dumps({"metadata": metadata, "courses": courses}, ensure_ascii=False), encoding="utf-8"
    )
    report = {
        **metadata,
        "invalid_schedules": errors,
        "orphan_subject_codes": orphan_codes,
        "orphan_schedule_section_ids": orphan_schedule_ids,
        "srs_expected_courses": 12456,
        "srs_count_confirmed": False,
        "missing_fields": ["department", "prerequisites", "location", "academic_year", "semester"],
        "student_enrollments_imported": False,
    }
    (destination / "workbook-validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with gzip.open(destination / "catalog.json.gz", "wt", encoding="utf-8") as target:
        json.dump({"metadata": metadata, "courses": courses}, target, ensure_ascii=False)
    print(json.dumps(metadata, ensure_ascii=False, indent=2))
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).with_name("data"))
    args = parser.parse_args()
    import_data(args.source, args.output_dir)
